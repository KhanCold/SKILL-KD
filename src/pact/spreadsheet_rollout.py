"""Codegen multi-rollout for SpreadsheetBench.

This module implements a codegen-style agent that generates Python code directly
(no tool-call), executes it, feeds back errors, and iterates. Aligned with
SkillOpt's codegen multi mode for fair comparison.
"""
from __future__ import annotations

import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any


CODEGEN_SYSTEM_PROMPT = """You are an expert Python programmer specializing in spreadsheet manipulation. You will be given a user instruction together with a preview of an input .xlsx file. Your job is to write a single self-contained Python script that reads the input file at the path stored in the variable INPUT_PATH, performs the requested manipulation, and saves the result to OUTPUT_PATH. The variables INPUT_PATH and OUTPUT_PATH are already defined — do NOT assign values to them. Use only the standard library, openpyxl, and pandas. Do not print anything. Do not use input(). Return ONLY the Python code inside a single ```python ... ``` fenced block.

## Common Workflow

1. **Explore** the input file: list sheets, inspect headers, check dimensions.
2. **Write Python code** that uses the pre-defined `INPUT_PATH` and `OUTPUT_PATH` variables.
3. **Verify** the output file was created and contains the expected values.
4. **Confirm** the target cells/range contain the expected values.

## Library Selection

| Use case | Library |
|----------|---------|
| Preserve formulas, formatting, named ranges | `openpyxl` |
| Bulk data transformation, aggregation, sorting | `pandas` → write back with `openpyxl` |
| Simple cell read/write | `openpyxl` |

**Warning**: `pandas.to_excel()` silently destroys existing formulas and named ranges.
When writing back to a spreadsheet that contains formulas, always use `openpyxl.save()`.

## solution.py Template

```python
import openpyxl
import pandas as pd

# INPUT_PATH and OUTPUT_PATH are already defined by the runtime
wb = openpyxl.load_workbook(INPUT_PATH)
ws = wb.active  # or wb["SheetName"]

# --- perform manipulation ---

wb.save(OUTPUT_PATH)
```

## Output Requirements

- Save the result to `OUTPUT_PATH`.
- Do not hardcode row counts or column letters — iterate over actual rows in the workbook.
- Preserve sheets and cells not mentioned in the instruction.

{skill_section}"""


def _build_system(skill_text: str) -> str:
    """Build system prompt with skill injection."""
    if skill_text.strip():
        skill_section = f"## Skill\n{skill_text.strip()}\n\n"
    else:
        skill_section = ""
    return CODEGEN_SYSTEM_PROMPT.format(skill_section=skill_section)


def _preview_workbook(path: str, max_rows: int = 5, max_cols: int = 20) -> str:
    """Generate a text preview of the first few rows of each sheet.

    Aligned with SkillOpt's _preview_workbook in codegen_agent.py.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    chunks: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        chunks.append(
            f"## Sheet: {sheet_name}  "
            f"(dim={ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column})"
        )
        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, max_rows),
            max_col=min(ws.max_column, max_cols),
        ):
            cells = []
            for cell in row:
                val = cell.value
                if val is None:
                    cells.append(f"{cell.coordinate}=")
                else:
                    val_str = str(val)
                    if len(val_str) > 40:
                        val_str = val_str[:37] + "..."
                    cells.append(f"{cell.coordinate}={val_str}")
            chunks.append(" | ".join(cells))
        if ws.max_row > max_rows:
            chunks.append(f"... ({ws.max_row - max_rows} more rows)")
        chunks.append("")  # Blank line between sheets, aligned with SkillOpt
    wb.close()
    return "\n".join(chunks)


def _build_user(
    instruction: str,
    instruction_type: str = "",
    answer_position: str = "",
    preview: str = "",
) -> str:
    """Build user prompt with task details.

    Aligned with SkillOpt's _build_user in codegen_agent.py.
    """
    parts: list[str] = []
    parts.append(f"# Instruction\n{instruction}")
    if instruction_type:
        parts.append(f"Instruction type: {instruction_type}")
    if answer_position:
        parts.append(f"Expected answer position: {answer_position}")
    if preview:
        parts.append(f"# Input spreadsheet preview\n{preview}")
    parts.append(
        "# Task\n"
        "Write a Python script that reads the workbook from the variable `INPUT_PATH`, "
        "applies the instruction, and writes the modified workbook to `OUTPUT_PATH`. "
        "Preserve all other cells unchanged. The preview may be truncated -- do not "
        "hardcode row counts or assume the data ends at the last previewed row; iterate "
        "over all actual rows in the workbook instead. Return only a ```python``` code block."
    )
    return "\n\n".join(parts)


def _extract_code(text: str) -> str:
    """Extract the first fenced Python code block from text.

    Aligned with SkillOpt's extract_code in codegen_agent.py.
    """
    if "```" not in text:
        return text.strip()
    start = text.find("```")
    nl = text.find("\n", start)
    end = text.find("```", nl + 1)
    if nl == -1 or end == -1:
        return text.strip()
    return text[nl + 1 : end].strip()


def _run_code(
    code: str,
    input_path: str,
    output_path: str,
    timeout: int = 120,
) -> tuple[bool, str]:
    """Execute generated code with path injection.

    Aligned with SkillOpt's run_generated_code in executor.py.
    Strips user-defined INPUT_PATH/OUTPUT_PATH, injects correct paths.
    Returns (success, error_output). Error output truncated to 5000 chars.
    """
    # Ensure output directory exists
    if output_path:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # Strip user-defined path assignments (aligned with SkillOpt's regex)
    code = re.sub(
        r'^\s*(INPUT_PATH|OUTPUT_PATH)\s*=\s*.+$',
        "",
        code,
        flags=re.MULTILINE,
    )

    # Build runner script
    runner = f"""
import sys
import traceback

INPUT_PATH = {repr(input_path)}
OUTPUT_PATH = {repr(output_path)}

try:
{chr(10).join('    ' + line for line in code.splitlines())}
except Exception:
    traceback.print_exc()
    sys.exit(2)
"""

    # Write to temp file and execute
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".py", delete=False, encoding="utf-8"
    ) as f:
        f.write(runner)
        tmp_path = f.name

    try:
        proc = subprocess.run(
            [sys.executable, tmp_path],
            capture_output=True,
            text=True,
            timeout=timeout,
        )
        success = proc.returncode == 0 and os.path.exists(output_path)
        if proc.returncode != 0:
            error = (proc.stdout + proc.stderr)[-5000:]
        elif not os.path.exists(output_path):
            error = "output file was not created"
        else:
            error = ""
        return success, error
    except subprocess.TimeoutExpired:
        return False, f"timeout after {timeout}s"
    except Exception as e:
        return False, str(e)[:5000]
    finally:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass


# ---------- Gold verification (ported from SkillOpt) ----------

_MAX_REPORT_CHARS = 12000


def _auto_verify_output(
    pred_path: str,
    gold_path: str,
    answer_position: str,
) -> str:
    """Compare predicted output with gold answer at cell level.

    Returns a human-readable verification report showing which cells are
    correct (✓) and which are wrong (✗). Used to provide cell-level feedback
    during training without leaking expected values.

    Ported from SkillOpt rollout.py:108-207.
    """
    import openpyxl
    from .spreadsheet_official import _compare_cell_value, _generate_cell_names

    if not os.path.exists(pred_path):
        return "Verification: output file does not exist."
    try:
        wb_pred = openpyxl.load_workbook(pred_path, data_only=True)
        wb_gold = openpyxl.load_workbook(gold_path, data_only=True)
    except Exception as e:
        return f"Verification: could not open workbooks: {e}"

    lines = ["## Output Verification"]
    try:
        for scr in (answer_position or "").split(","):
            scr = scr.strip()
            if not scr:
                continue
            if "!" in scr:
                sheet_name, cell_range = scr.split("!", 1)
                sheet_name = sheet_name.strip().strip("'\"")
            else:
                sheet_name = wb_gold.sheetnames[0]
                cell_range = scr
            cell_range = cell_range.strip().strip("'\"")

            cell_names = _generate_cell_names(cell_range)
            ws_pred = wb_pred[sheet_name] if sheet_name in wb_pred.sheetnames else None
            ws_gold = wb_gold[sheet_name] if sheet_name in wb_gold.sheetnames else None

            if ws_pred is None:
                lines.append(f"  Sheet '{sheet_name}' NOT FOUND in output.")
                continue

            n_empty_correct = 0
            for cn in cell_names:
                gv = ws_gold[cn].value if ws_gold else "N/A"
                pv = ws_pred[cn].value
                ok_cell = ws_gold is not None and _compare_cell_value(gv, pv)
                # Collapse cells that are correct AND empty on both sides
                if ok_cell and gv in (None, "") and pv in (None, ""):
                    n_empty_correct += 1
                    continue
                match = "✓" if ok_cell else "✗"
                lines.append(f"  {sheet_name}!{cn}: got={pv!r}, expected={gv!r} {match}")
            if n_empty_correct:
                lines.append(f"  (+{n_empty_correct} empty cells correct, omitted)")

        # Check for formula strings in output
        formula_cells = []
        for sn in wb_pred.sheetnames:
            ws = wb_pred[sn]
            for row in ws.iter_rows(max_row=min(ws.max_row, 200), values_only=False):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_cells.append(f"{sn}!{cell.coordinate}={cell.value}")
                        if len(formula_cells) >= 10:
                            break
                if len(formula_cells) >= 10:
                    break
            if len(formula_cells) >= 10:
                break
        if formula_cells:
            lines.append(f"\n  WARNING: {len(formula_cells)} cells contain Excel formulas (openpyxl cannot evaluate them):")
            for fc in formula_cells[:5]:
                lines.append(f"    {fc}")
            if len(formula_cells) > 5:
                lines.append(f"    ... and {len(formula_cells) - 5} more")
    finally:
        wb_pred.close()
        wb_gold.close()

    report = "\n".join(lines)
    # Head+tail truncation for oversized reports
    if len(report) > _MAX_REPORT_CHARS:
        half = _MAX_REPORT_CHARS // 2
        report = (
            report[:half]
            + f"\n  ...[verification report truncated, {len(report)} chars total]...\n"
            + report[-half:]
        )
    return report


def _build_eval_feedback(verify_report: str) -> str:
    """Build feedback from verification report without leaking expected values.

    Strips the ``expected=...`` part so the model sees only its own output
    and whether each cell is correct or wrong.

    Ported from SkillOpt codegen_agent.py:46-83.
    """
    wrong_lines = []
    n_correct = 0
    for raw_line in verify_report.splitlines():
        raw_line = raw_line.strip()
        if not raw_line:
            continue
        # Match lines like "Sheet1!D2: got=None, expected=0 ✗"
        m = re.match(
            r"(\S+!?\w+):\s*got=(.+?),\s*expected=.+?\s*(✓|✗)$",
            raw_line,
        )
        if m:
            cell, got_val, mark = m.groups()
            if mark == "✗":
                wrong_lines.append(f"  {cell}: your output = {got_val} (WRONG)")
            else:
                n_correct += 1
    lines = ["Your code executed successfully but produced incorrect results.",
             "The following cells have wrong values:"]
    lines.extend(wrong_lines)
    if n_correct:
        lines.append(f"  ({n_correct} other cells are correct.)")
    lines.append(
        "\nPlease analyze the spreadsheet data more carefully and fix the code. "
        "Return a complete corrected Python script inside a ```python``` block."
    )
    return "\n".join(lines)


def run_codegen_multi(
    task: dict[str, Any],
    skill_text: str,
    llm: Any,
    role: str = "student",
    output_dir: Path | None = None,
    max_turns: int = 5,
    max_completion_tokens: int = 16384,
    gold_path: str = "",
) -> dict[str, Any]:
    """Run codegen multi-round agent for a SpreadsheetBench task.

    Generates Python code, executes it, feeds back errors, iterates up to
    max_turns rounds. Aligned with SkillOpt's run_multi in codegen_agent.py.

    When *gold_path* is provided (training/evolve mode), a successful
    execution is followed by a cell-level verification against the gold
    answer.  If the output is wrong the model receives per-cell feedback
    (without leaking expected values) and gets another turn.  When
    *gold_path* is empty (eval/test mode) execution success is sufficient
    to stop.

    Args:
        task: Task dict with instruction, spreadsheet_path, id, etc.
        skill_text: Skill rules to inject into system prompt.
        llm: LLMClient instance.
        role: Role label for trajectory (student/teacher/student_retry).
        output_dir: Directory to save artifacts.
        max_turns: Maximum number of LLM calls.
        max_completion_tokens: Max tokens for each LLM response.
        gold_path: Path to golden answer xlsx for eval feedback during
            training.  Leave empty for eval/test to avoid data leakage.

    Returns:
        Dict with role, skill_used, prompt, raw, final_answer, conversation, n_turns.
    """
    # Extract task metadata
    instruction = task.get("instruction", "")
    task_id = str(task.get("id", ""))
    dataset_root = Path(task.get("dataset_root", ""))
    spreadsheet_path = task.get("spreadsheet_path", f"spreadsheet/{task_id}")
    instruction_type = task.get("instruction_type", "")
    answer_position = task.get("answer_position", "")

    # Build paths (aligned with current ReAct implementation)
    spreadsheet_dir = dataset_root / spreadsheet_path
    input_file = f"1_{task_id}_input.xlsx"
    input_path = spreadsheet_dir / input_file
    output_file = f"1_{task_id}_output.xlsx"

    # Create isolated work directory
    work_dir = tempfile.mkdtemp(prefix=f"pact_codegen_{task_id}_")
    output_path = Path(work_dir) / output_file

    try:
        # Build prompts
        preview = _preview_workbook(str(input_path)) if input_path.exists() else ""
        system = _build_system(skill_text)
        user = _build_user(
            instruction,
            instruction_type=instruction_type,
            answer_position=answer_position,
            preview=preview,
        )

        # Initialize message history
        messages = [
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ]

        # Iteration loop
        code = ""
        raw = ""
        conversation: list[dict[str, str]] = []
        n_turns = 0

        for turn in range(max_turns):
            n_turns = turn + 1

            # Call LLM (no tools)
            response = llm.chat_raw(
                messages,
                max_completion_tokens=max_completion_tokens,
            )
            raw = response.content or ""
            conversation.append({"role": "assistant", "content": raw})

            # Extract code
            code = _extract_code(raw)
            if not code:
                # No code found, ask for code
                feedback = (
                    "No Python code block was found in your response. "
                    "Please return a complete Python script inside a ```python``` block."
                )
                messages.append({"role": "assistant", "content": raw})
                messages.append({"role": "user", "content": feedback})
                conversation.append({"role": "user", "content": feedback})
                continue

            # Execute code
            success, error = _run_code(code, str(input_path), str(output_path))
            if success:
                if gold_path and answer_position:
                    # Training mode: verify output against gold answer
                    from .spreadsheet_official import evaluate as _ssb_evaluate
                    eval_result = _ssb_evaluate(
                        str(output_path), gold_path, instruction_type, answer_position,
                    )
                    if eval_result["ok"]:
                        break  # Correct answer — stop
                    # Wrong answer: build cell-level feedback (no golden leakage)
                    verify = _auto_verify_output(str(output_path), gold_path, answer_position)
                    feedback = _build_eval_feedback(verify)
                else:
                    # Eval/test mode: execution success is sufficient
                    break

            if success:
                # Gold verification failed — feedback already built above
                pass
            else:
                # Execution failed — feed back the traceback
                feedback = (
                    f"The code raised an error during execution:\n\n"
                    f"```\n{error}\n```\n\n"
                    f"Please fix the code and return a complete corrected Python script "
                    f"inside a ```python``` block."
                )
            messages.append({"role": "assistant", "content": raw})
            messages.append({"role": "user", "content": feedback})
            conversation.append({"role": "user", "content": feedback})

        # Build return trajectory
        prompt_text = f"<system>\n{system}\n\n<user>\n{user}"
        raw_with_code = f"```python\n{code}\n```" if code else raw

        return {
            "role": role,
            "skill_used": skill_text,
            "prompt": prompt_text,
            "raw": raw_with_code,  # For evaluate() to extract_code()
            "final_answer": code,
            "conversation": conversation,
            "n_turns": n_turns,
        }

    finally:
        # Clean up work directory
        import shutil
        shutil.rmtree(work_dir, ignore_errors=True)
